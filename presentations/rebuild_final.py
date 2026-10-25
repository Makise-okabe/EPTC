"""Judge-facing EPTC presentation revision.

Builds on the editable 30-slide deck. Adds the previously missing 45-W
verification motivation, original material-mechanics evidence, rigorous source
provenance, and matched strain data, then arranges all slides in the actual
research progression. The source deck must be built first.

IMPORTANT: The source deck's header/footer emulate the organisers' sample.
If an original native organiser template is supplied to the GitHub repository,
reuse that native master in a further version; do not claim an emulation is
the exact source master.
"""
from pathlib import Path
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from deck_helpers import *

ROOT=Path(__file__).resolve().parent
F=ROOT/"EPTC_2026_Student_Innovation_Challenge_Presentation.pptx"
p=Presentation(str(F))
assert len(p.slides)==30

def citation(sl,msg):
    # Replace the smaller recurring source line with a citable precise source;
    # keep the unchanged conference header and footer.
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        if sh.top/Inches(1) < 6.68 or sh.top/Inches(1) > 7.07: continue
        if ("Source:" in sh.text or "Literature:" in sh.text
            or "References:" in sh.text or "Geometry:" in sh.text
            or "Published TIM context:" in sh.text
            or "Coverage fractions" in sh.text
            or "Physical interpretation:" in sh.text
            or "Project 700-W simulation" in sh.text):
            sh.text=msg
            sh.left=Inches(.56);sh.width=Inches(12.17)
            sh.height=Inches(.29)
            for pa in sh.text_frame.paragraphs:
                for ru in pa.runs:
                    ru.font.name="Aptos";ru.font.size=Pt(8.4)
                    ru.font.color.rgb=rgb(GREY)
            return True
    # Some original slides did not contain a dedicated citation text box.
    # Insert a clean one in the reserved source strip rather than aborting.
    put_rect(sl,.53,6.74,12.24,.29,fill=WHITE)
    put_text(sl,.56,6.77,12.16,.23,msg,size=8.3,color=GREY)
    return True

# Full identifiers for previously implicit literature-derived claims.
doi_wen="Wen et al., Int. J. Therm. Sci. 221 (2026), 110432. DOI: 10.1016/j.ijthermalsci.2025.110432."
doi_yi="Yi et al., Microelectron. Reliab. 166 (2025), 115607. DOI: 10.1016/j.microrel.2025.115607."
doi_li="Li et al., Compos. Sci. Technol. 234 (2023), 109948. DOI: 10.1016/j.compscitech.2023.109948."
doi_zhao="Zhao et al., Int. J. Heat Mass Transfer 140 (2019), 705-716. DOI: 10.1016/j.ijheatmasstransfer.2019.06.045."

assert citation(p.slides[5],"Copper lid / TIM2 / substrate: "+doi_wen+" Die, cooler thickness and 700 W: study-defined.")
assert citation(p.slides[6],"TIM conductivity / coverage: "+doi_yi+" Geometry anchors: "+doi_wen)
assert citation(p.slides[8],"Thermal contact formulation: Prasher (2006), DOI: 10.1109/JPROC.2006.879796; "+doi_zhao)
assert citation(p.slides[11],"Published retained-area fractions only: "+doi_yi+" Thermal temperatures: team's simulations.")
assert citation(p.slides[12],"Coverage fractions: "+doi_yi+" Spatial loss maps and all temperatures: study-defined simulations.")
assert citation(p.slides[13],"Team's nine Icepak simulations; parameter plausibility: "+doi_li)
assert citation(p.slides[15],"Team's recovered four-case Icepak results; retained area motivated by "+doi_yi)
assert citation(p.slides[25],"Reliability model limitations: Subramanian et al. (2019), DOI: 10.1115/1.4042805; "+doi_zhao)
assert citation(p.slides[26],"Qualification approach: Due & Robinson (2013), DOI: 10.1016/j.applthermaleng.2012.06.013; "+doi_zhao)

# Emphasize precise "baseline validation" rather than ambiguous proof of physics.
for sh in p.slides[9].shapes:
    if sh.has_text_frame and "They support solver consistency, not experimental validation." in sh.text:
        sh.text=("These are numerical consistency checks—not experimental calibration. "
                 "The later 700-W study assumes ideal thermal contact.")
        for pa in sh.text_frame.paragraphs:
            for ru in pa.runs:ru.font.name="Aptos";ru.font.size=Pt(18);ru.font.color.rgb=rgb(INK)

# 31: Why start with 45 W?
s=addslide(p,"Why begin with a 45-W analytical model?","RESEARCH METHOD / FOUNDATION",
           "Source: team's 45-W X23 verification and mesh logs; the 0.200 °C/W reference was study-selected.")
put_text(s,.70,1.28,11.92,.53,
         "First establish a controlled heat path before simulating a 700-W accelerator.",
         25,NAVY,True)
items=[
 ("Analytical expectation","Rbulk = 0.115385 °C/W","45 W through the reference X23 TIM gives ~5.192 °C bulk drop.",TEAL),
 ("Numerical convergence","0.000387%","Medium-to-fine change in die peak temperature; mesh refinement alone was stable.",GREEN),
 ("Calibration limitation","27.72%","FEA reference resistance 0.144568 °C/W versus selected 0.200 °C/W comparison target.",RUST)
]
for i,(name,value,detail,c) in enumerate(items):
    x=.62+i*4.23
    card(s,x,2.07,3.96,3.15,fill=PALE)
    put_rect(s,x+.20,2.28,.10,.62,fill=c)
    put_text(s,x+.41,2.33,3.27,.43,name,17,NAVY,True)
    put_text(s,x+.40,3.04,3.22,.65,value,28,c,True)
    put_text(s,x+.41,3.90,3.18,1.05,detail,14,INK)
card(s,.82,5.52,11.66,.86,fill=PALE_BLUE)
put_text(s,1.07,5.74,11.14,.48,
         "Why it mattered: detect model/interface assumptions before comparing materials under 700 W.",
         18,NAVY,True,align=PP_ALIGN.CENTER)

# 32: What the initial mismatch means (NOT arbitrary contact calibration).
s=addslide(p,"What the initial verification did—and did not—establish",
           "RESEARCH METHOD / INTERFACE VALIDITY",
           "Source: team's separate 45-W analytical, mesh and calibration diagnostic; not a physical contact measurement.")
card(s,.63,1.49,6.05,3.77,fill=PALE)
put_text(s,.94,1.82,5.37,.41,"Model observation",20,NAVY,True)
bullets(s,.96,2.52,5.43,[
    "Analytical layer resistance and FEA package resistance are different quantities.",
    "The mesh check was stable; agreement with an external target was not established.",
    "The mismatch motivated separating bulk and interface contributions."
],size=16,gap=.79)
card(s,6.88,1.49,5.80,3.77,fill=LIGHT)
put_text(s,7.16,1.82,5.20,.41,"Scope of the subsequent study",20,NAVY,True)
bullets(s,7.21,2.51,5.08,[
    "An 8.65 °C·mm²/W effective fitting term was explored in a separate diagnostic.",
    "It was NOT transferred to the 700-W material and DOE simulations.",
    "Consequently, the main comparisons use ideal thermal interfaces."
],size=16,gap=.79)
card(s,.78,5.55,11.85,.78,fill=PALE_BLUE)
put_text(s,1.06,5.76,11.21,.39,
         "Interpret the subsequent temperatures as model-specific comparative results, not qualified junction temperatures.",
         17,RUST,True,align=PP_ALIGN.CENTER)

# 33: Early grease and indium mechanics: explain why thermal material ranking is not reliability ranking.
s=addslide(p,"Why a thermal material screen was not enough",
           "RESEARCH PROGRESSION / INITIAL MECHANICS",
           "Team's MAPDL grease/indium results; indium constitutive sensitivity is a study assumption, not fatigue prediction.")
put_text(s,.72,1.30,11.88,.47,"Thermal benefit did not establish the mechanically preferable TIM.",23,NAVY,True)
card(s,.68,2.03,3.91,3.38,fill=PALE)
put_text(s,.94,2.31,3.30,.35,"COMPLIANT GREASE",15,TEAL,True)
put_text(s,.93,2.86,3.33,.62,"0.00105 MPa",25,TEAL,True)
put_text(s,.93,3.55,3.37,.90,"Mean modeled TIM stress; assumed E = 0.05 MPa. Warpage ≈ 15.51 µm.",15,INK)
card(s,4.74,2.03,3.91,3.38,fill=PALE)
put_text(s,5.00,2.31,3.30,.35,"ELASTIC INDIUM",15,RUST,True)
put_text(s,5.00,2.86,3.30,.62,"50.90 MPa",25,RUST,True)
put_text(s,5.00,3.55,3.36,.90,"Elastic mean stress is not credible as a realistic post-yield service prediction.",15,INK)
card(s,8.81,2.03,3.82,3.38,fill=PALE)
put_text(s,9.06,2.31,3.35,.35,"PLASTICITY SENSITIVITY",14,GREEN,True)
put_text(s,9.06,2.86,3.34,.62,"0.51% / 0.49%",21,GREEN,True)
put_text(s,9.06,3.55,3.34,1.09,"Mean plastic strain at assumed 1.00 / 2.14 MPa yields; no lifecycle prediction.",14.5,INK)
put_text(s,1.07,5.76,11.24,.52,
         "Reason for the next step: use carbon design variables and a consistent mechanical model.",
         18,NAVY,True,align=PP_ALIGN.CENTER)

# 34: Source-to-model crosswalk (include unambiguous citations on dimensions).
s=addslide(p,"Dimensions and literature: what came from where",
           "METHODS / PARAMETER PROVENANCE",
           doi_wen+" "+doi_yi)
columns=[.53,3.12,8.09]; widths=[2.45,4.87,4.61]
for x,w,name in zip(columns,widths,["Package element","Modeled input","Source classification"]):
    put_rect(s,x,1.46,w,.56,fill=NAVY)
    put_text(s,x+.09,1.61,w-.14,.32,name,14,WHITE,True)
rows=[
 ("Copper lid","52.4 × 52.4 × 0.70 mm","Wen et al. (2026), cited DOI"),
 ("TIM2","0.10 mm; 5.4 W/(m·K)","Wen et al. (2026), cited DOI"),
 ("Equivalent substrate","52.4 × 52.4 × 1.0 mm","Literature-aligned, Wen et al."),
 ("Carbon thermal anchor","21 W/(m·K); 88.22 / 64.82% area","Yi et al. (2025), cited DOI"),
 ("Die + imposed power","28.5 mm die; 700 W, 4 × 4 map","Study-defined—not measured"),
 ("Cold plate / boundary","60 × 60 × 8 mm; fixed 25 °C","Footprint literature-aligned; rest assumed"),
 ("Design variables","k = 15/21/30; BLT = 100/250/400 µm","Simulation DOE, not a vendor specification")
]
for i,row in enumerate(rows):
    y=2.10+i*.62
    fill=WHITE if i%2 else PALE
    for x,w,t in zip(columns,widths,row):
        put_rect(s,x,y,w,.59,fill=fill,line=LINE)
        put_text(s,x+.09,y+.10,w-.16,.42,t,12.7,NAVY if x==columns[0] else INK,
                 bold=(x==columns[0]))
put_text(s,.72,6.69,12.10,.19,
         "DOIs: Wen 10.1016/j.ijthermalsci.2025.110432; Yi 10.1016/j.microrel.2025.115607",
         9.4,GREY)

# 35: All-strain trends complement the stress and warpage plots.
s=addslide(p,"Mean TIM strain confirms the intermediate-BLT minimum",
           "RESULTS / MATCHED MECHANICS",
           "Team's nine matching Icepak-to-MAPDL MID-model simulations (dimensionless strain shown as %).")
card(s,.65,1.40,8.72,4.77)
xy_plot(s,.90,1.82,8.15,3.94,[
 ("k = 15 W/(m·K)",[1.2568,1.2296,1.3535]),
 ("k = 21 W/(m·K)",[1.2400,1.1594,1.2369]),
 ("k = 30 W/(m·K)",[1.2279,1.1076,1.1498])
],[100,250,400],yrange=(1.08,1.38),xlabels=["100 µm","250 µm","400 µm"],yfmt="{:.2f}")
card(s,9.61,1.40,3.00,4.77,fill=LIGHT)
put_text(s,9.87,1.78,2.47,.45,"A matched result",17,NAVY,True)
put_text(s,9.87,2.50,2.45,.60,"1.1076%",26,GREEN,True)
put_text(s,9.88,3.24,2.43,1.36,
         "Lowest sampled equivalent elastic strain at k = 30 W/(m·K), BLT = 250 µm.",
         15,INK)
put_text(s,1.01,6.27,11.31,.39,
         "At every sampled conductivity, the 250-µm case minimizes mean stress AND mean strain.",
         17,NAVY,True,align=PP_ALIGN.CENTER)

assert len(p.slides)==35, len(p.slides)

# Chronological evidence progression. Original slides indexed 1..30; new 31..35.
new_order=[
 1,3,4,5,                  # challenge, objectives, methodology
 6,34,7,9,                # geometry / dimensional provenance / assumptions / physics
 31,32,                   # 45-W reason and calibration boundary
 8,10,11,33,              # 700-W workload, numerical checks, material screen, early mechanics
 12,13,                   # coverage, topology
 14,15,16,                # thermal DOE, trend, retained-contact verification
 17,18,19,                # one-way handoff, property bracket, mesh audit
 20,21,35,22,23,24,       # matched design, stress, strain, warpage, fixed-k contrast, decision map
 25,2,                    # recommendation followed by compact executive synthesis
 26,27,28,                # limitations, qualification route, takeaways
 29,30                    # references and backup full dataset
]
assert len(new_order)==35 and sorted(new_order)==list(range(1,36))
ids=list(p.slides._sldIdLst)
for el in list(p.slides._sldIdLst):p.slides._sldIdLst.remove(el)
for ix in new_order:p.slides._sldIdLst.append(ids[ix-1])

# Correct the footer numbering after reordering, retain the conference footer style.
for n,slide in enumerate(p.slides,1):
    for sh in slide.shapes:
        if not sh.has_text_frame:continue
        x=sh.left/Inches(1);y=sh.top/Inches(1)
        if x>=12.34 and y>=7.07 and sh.text.strip().isdigit():
            sh.text=str(n)
            for par in sh.text_frame.paragraphs:
                par.alignment=PP_ALIGN.RIGHT
                for run in par.runs:
                    run.font.name="Aptos";run.font.size=Pt(12.5);run.font.color.rgb=rgb(WHITE)
    # Presenter readability: coherent modern fonts, 8.2pt minimum source text.
    for sh in slide.shapes:
        if not sh.has_text_frame:continue
        y=sh.top/Inches(1)
        for pa in sh.text_frame.paragraphs:
            for run in pa.runs:
                run.font.name="Aptos Display" if y<.78 else "Aptos"
                if run.font.size and run.font.size.pt<8.2:run.font.size=Pt(8.2)

# Persist and audit original numeric results, citations and structure.
all_text="\n".join(sh.text for sl in p.slides for sh in sl.shapes if sh.has_text_frame)
for token in ["0.000387%","0.144568","0.200 °C/W","50.90 MPa",
              "0.010841","9.46%","80.650","64.82%",
              "10.1016/j.microrel.2025.115607",
              "10.1016/j.ijthermalsci.2025.110432"]:
    assert token in all_text,token
for banned in ["V6A","V4B","V4D","V3B","Draft"]:
    assert banned not in all_text,banned
issues=[]
for si,sl in enumerate(p.slides,1):
    for sh in sl.shapes:
        x=sh.left/Inches(1);y=sh.top/Inches(1)
        w=sh.width/Inches(1);h=sh.height/Inches(1)
        if x<-.05 or y<-.05 or x+w>13.333+.12 or y+h>7.5+.12:
            issues.append((si,round(x,3),round(y,3),round(w,3),round(h,3)))
if issues:raise AssertionError("Off-slide shapes "+str(issues[:12]))
p.save(F)
checked=Presentation(F)
assert len(checked.slides)==35
print("PRESENTATION REVISED",F)
print("SLIDES",len(checked.slides),"BYTES",F.stat().st_size)
print("NUMERICAL + REFERENCE + LAYOUT QA PASS")
