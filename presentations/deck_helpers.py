"""Editable PowerPoint primitives for the EPTC 2026 student competition deck."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
import math

W,H=13.333,7.50
TOP=0.78
BOTTOM=0.39
CONTENT_TOP=1.04
CONTENT_BOTTOM=7.03

RUST="963A21"
RUST_DARK="702717"
CORAL="BC6145"
NAVY="183A52"
NAVY2="294B61"
TEAL="13868A"
GREEN="2A896A"
ORANGE="E29E40"
RED="BB5240"
GREY="657585"
INK="203344"
LIGHT="F0F5F7"
PALE="F8FAFB"
PALE_BLUE="E5EFF3"
WHITE="FFFFFF"
LINE="D9E3E8"

def rgb(h):
    h=h.strip("#")
    return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))

def put_rect(sl,x,y,w,h,fill=WHITE,line=None,radius=False,width=0.65):
    shape=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                              Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill)
    if line:
        shape.line.color.rgb=rgb(line);shape.line.width=Pt(width)
    else: shape.line.fill.background()
    return shape

def put_line(sl,x1,y1,x2,y2,color=LINE,width=1.2):
    shape=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    shape.line.color.rgb=rgb(color);shape.line.width=Pt(width)
    return shape

def put_text(sl,x,y,w,h,text,size=18,color=INK,bold=False,align=None,valign=None,
             margin=0.015,font="Aptos",italic=False):
    shape=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin)
    tf.margin_top=tf.margin_bottom=Inches(0.01)
    if valign is not None:tf.vertical_anchor=valign
    for k,line in enumerate(str(text).split("\n")):
        p=tf.paragraphs[0] if k==0 else tf.add_paragraph()
        p.text=line;p.space_before=Pt(0);p.space_after=Pt(0)
        if align is not None:p.alignment=align
        for run in p.runs:
            run.font.name=font;run.font.size=Pt(size)
            run.font.color.rgb=rgb(color);run.font.bold=bold;run.font.italic=italic
    return shape

def pill(sl,x,y,text,color=TEAL,w=None,h=0.37,small=12):
    if w is None:w=max(1.15,min(5.4,0.105*len(text)+0.35))
    put_rect(sl,x,y,w,h,fill=color,radius=True)
    put_text(sl,x+.10,y+.065,w-.20,h-.11,text,size=small,bold=True,color=WHITE,
             align=PP_ALIGN.CENTER)
    return w

def label(sl,x,y,text,color=TEAL,size=11):
    put_text(sl,x,y,9,.24,text.upper(),size=size,color=color,bold=True)

def card(sl,x,y,w,h,title=None,fill=WHITE,border=LINE,title_color=NAVY):
    put_rect(sl,x,y,w,h,fill=fill,line=border,radius=True,width=.85)
    if title:put_text(sl,x+.20,y+.16,w-.40,.37,title,size=16,bold=True,color=title_color)
    return (x+.19,y+.18,w-.38,h-.36)

def bullets(sl,x,y,w,items,size=18,gap=.53,color=INK,dot=TEAL):
    for i,txt in enumerate(items):
        yy=y+i*gap
        put_rect(sl,x,yy+.13,.085,.085,fill=dot,radius=True)
        put_text(sl,x+.19,yy,w-.23,gap-.03,txt,size=size,color=color)

def stat(sl,x,y,w,h,value,descriptor,color=TEAL,detail=None):
    card(sl,x,y,w,h)
    put_text(sl,x+.15,y+.22,w-.30,.57,value,size=31,color=color,bold=True)
    put_text(sl,x+.15,y+.91,w-.30,.61,descriptor,size=15,color=INK,bold=True)
    if detail:put_text(sl,x+.15,y+h-.48,w-.30,.28,detail,size=11,color=GREY)

def textbox_card(sl,x,y,w,h,title,desc,accent=TEAL,size=16):
    card(sl,x,y,w,h)
    put_rect(sl,x+.18,y+.20,.08,h-.40,fill=accent)
    put_text(sl,x+.39,y+.17,w-.60,.45,title,size=17,color=NAVY,bold=True)
    put_text(sl,x+.39,y+.67,w-.65,h-.79,desc,size=size,color=INK)

def branded(sl,title,num,section=None,source=None):
    """Recreates the visual language of the supplied organiser slide:
    rust PCB-inspired top/bottom bands, EPTC mark, central white title."""
    put_rect(sl,0,0,W,TOP,fill=RUST)
    put_rect(sl,0,0,1.85,TOP,fill=RUST_DARK)
    for j in range(10):
        ax=8.15+j*.57
        put_line(sl,ax,.13,ax+.29,.68,color=CORAL,width=.65)
        put_rect(sl,ax+.24,.62,.055,.055,fill="C87851")
    for j in range(6):
        put_line(sl,10.0+j*.45,.06,10.85+j*.45,.56,color="B05C3D",width=.42)
    put_text(sl,.23,.11,1.45,.35,"EPTC",size=25,color=WHITE,bold=True,font="Arial")
    put_text(sl,1.07,.43,.58,.20,"2026",size=12,color=WHITE,bold=True)
    put_line(sl,.24,.59,1.62,.59,color=WHITE,width=.9)
    size=27 if len(title)<52 else 23 if len(title)<70 else 20
    put_text(sl,1.94,.16,10.99,.45,title,size=size,color=WHITE,bold=True,
             align=PP_ALIGN.CENTER)
    put_rect(sl,0,H-BOTTOM,W,BOTTOM,fill=RUST_DARK)
    for j in range(16):
        put_line(sl,7.5+j*.34,H-.34,7.72+j*.34,H-.045,color=RUST,width=.6)
    put_text(sl,2.2,H-.315,9.0,.21,
             "28th Electronics Packaging Technology Conference  |  1–4 Dec 2026  |  Singapore",
             size=10.5,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
    put_text(sl,12.45,H-.325,.48,.27,str(num),size=13,color=WHITE,align=PP_ALIGN.RIGHT)
    if section:label(sl,.58,.91,section,color=CORAL,size=9)
    if source:
        put_text(sl,.55,6.86,12.10,.15,source,size=8.1,color=GREY)

def addslide(prs,title,section=None,source=None,notes=None):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    put_rect(slide,0,TOP,W,H-TOP-BOTTOM,fill=WHITE)
    branded(slide,title,len(prs.slides),section,source)
    if notes:
        try:
            tf=slide.notes_slide.notes_text_frame
            if tf:tf.text=notes
        except Exception:pass
    return slide

def bar_h(sl,x,y,w,h,labels,values,colors=None,vmin=0,vmax=None,fmt="{:.1f}",
          unit="",fs=14,value_fs=14,label_width=2.4):
    if vmax is None:vmax=max(values)*1.12
    x0=x+label_width; barw=w-label_width-.72
    row_h=h/max(len(labels),1)
    for i,(k,v) in enumerate(zip(labels,values)):
        yy=y+i*row_h
        put_text(sl,x,yy+.12,label_width-.12,row_h-.22,k,size=fs,color=INK,bold=(i==0))
        ww=max(.03,(v-vmin)/(vmax-vmin)*barw)
        put_rect(sl,x0,yy+.14,barw,.35,fill=LIGHT,radius=True)
        put_rect(sl,x0,yy+.14,ww,.35,fill=(colors[i] if colors else TEAL),radius=True)
        put_text(sl,x0+min(ww+.08,barw+.03),yy+.15,.70,.31,fmt.format(v)+unit,size=value_fs,color=NAVY,bold=True)

def bar_v(sl,x,y,w,h,labels,vals,colors=None,vmin=0,vmax=None,unit="",value_dec=1,fs=12):
    vmax=vmax if vmax is not None else max(vals)*1.10
    plot_h=h-.70
    plot_y=y+plot_h
    put_line(sl,x+.08,plot_y,x+w-.03,plot_y,color=GREY,width=.9)
    n=len(vals)
    for i,(lab,v) in enumerate(zip(labels,vals)):
        cw=w/n
        cx=x+i*cw+cw*.20
        bw=cw*.55
        scaled=(v-vmin)/(vmax-vmin)*plot_h*.94
        put_rect(sl,cx,plot_y-scaled,bw,scaled,fill=colors[i] if colors else TEAL,radius=False)
        put_text(sl,x+i*cw+.02,plot_y+.12,cw-.04,.36,lab,size=fs,color=INK,align=PP_ALIGN.CENTER)
        put_text(sl,cx-.12,plot_y-scaled-.33,bw+.24,.28,
                 (("%."+str(value_dec)+"f")%v)+unit,size=12,bold=True,color=NAVY,
                 align=PP_ALIGN.CENTER)

def colorinterp(a,b,t):
    aa=[int(a[k:k+2],16) for k in [0,2,4]]
    bb=[int(b[k:k+2],16) for k in [0,2,4]]
    return "".join("%02X"%int(aa[i]*(1-t)+bb[i]*t) for i in range(3))

def heatmap(sl,x,y,w,h,rows,cols,matrix,fmt="{:.3f}",low=None,high=None,palette=("E3F1ED","D55A3A"),value_size=17):
    mn=min(z for row in matrix for z in row) if low is None else low
    mx=max(z for row in matrix for z in row) if high is None else high
    left=1.09;top=.55
    sw=(w-left)/len(cols);sh=(h-top)/len(rows)
    for j,col in enumerate(cols):
        put_text(sl,x+left+j*sw,y+.01,sw,.33,col,size=13,color=NAVY,bold=True,align=PP_ALIGN.CENTER)
    for i,row in enumerate(rows):
        put_text(sl,x,y+top+i*sh+.28,left-.12,.40,row,size=13,color=NAVY,bold=True)
        for j,v in enumerate(matrix[i]):
            c=colorinterp(palette[0],palette[1],(v-mn)/(mx-mn) if mx>mn else 0)
            bx=x+left+j*sw+.026;by=y+top+i*sh+.026
            put_rect(sl,bx,by,sw-.06,sh-.065,fill=c,radius=True)
            put_text(sl,bx+.015,by+sh*.26,sw-.09,.39,fmt.format(v),
                     size=value_size,color=NAVY,bold=True,align=PP_ALIGN.CENTER)
    return mn,mx

def xy_plot(sl,x,y,w,h,series,xvals,yrange=None,colors=None,legend=True,
            labels=None,yfmt="{:.2f}",xlabels=None,mark=True):
    """Editable line chart, x values equally spaced; series=list of (label,[y...])."""
    colors=colors or [CORAL,TEAL,GREEN]
    allv=[v for _,vals in series for v in vals]
    lo=min(allv) if yrange is None else yrange[0]
    hi=max(allv) if yrange is None else yrange[1]
    if hi==lo:hi+=1
    L=.60;R=.22;T=.12;B=.74
    plx=x+L;ply=y+T;pw=w-L-R;ph=h-T-B
    for k in range(5):
        yy=ply+ph-k*ph/4
        put_line(sl,plx,yy,plx+pw,yy,color=LINE,width=.7)
        put_text(sl,x-.03,yy-.11,.57,.27,yfmt.format(lo+(hi-lo)*k/4),size=9.7,color=GREY,align=PP_ALIGN.RIGHT)
    pts={}
    for i,xv in enumerate(xvals):
        xx=plx+i*pw/(len(xvals)-1)
        put_text(sl,xx-.35,ply+ph+.12,.75,.25,str(xlabels[i] if xlabels else xv),
                 size=11,color=INK,align=PP_ALIGN.CENTER)
    for si,(serieslab,ys) in enumerate(series):
        cl=colors[si%len(colors)];coords=[]
        for k,val in enumerate(ys):
            px=plx+k*pw/(len(ys)-1)
            py=ply+ph-(val-lo)/(hi-lo)*ph
            coords.append((px,py))
            if mark:put_rect(sl,px-.045,py-.045,.09,.09,fill=cl,radius=True)
        for k in range(len(coords)-1):
            put_line(sl,*coords[k],*coords[k+1],color=cl,width=2.2)
        pts[serieslab]=coords
        if legend:
            lx=plx+.08+si*min(2.35,pw/len(series))
            ly=y+h-.32
            put_rect(sl,lx,ly,.16,.10,fill=cl)
            put_text(sl,lx+.24,ly-.08,2.10,.30,serieslab,size=10,color=INK)
    return pts

def scatter(sl,x,y,w,h,points,xrange,yrange,label_xy=None):
    """points: (x,y,cat,text), editable scatter plot."""
    colors={"15":CORAL,"21":TEAL,"30":GREEN}
    L=.76;R=.22;T=.17;B=.62
    px=x+L;py=y+T;pw=w-L-R;ph=h-T-B
    for k in range(5):
        yy=py+ph-k*ph/4
        put_line(sl,px,yy,px+pw,yy,LINE,.6)
        put_text(sl,x,yy-.11,.65,.20,"%.4f"%(yrange[0]+(yrange[1]-yrange[0])*k/4),
                 10,GREY,align=PP_ALIGN.RIGHT)
    for k in range(5):
        xx=px+k*pw/4
        put_line(sl,xx,py,xx,py+ph,LINE,.55)
        put_text(sl,xx-.33,py+ph+.14,.70,.22,"%.0f"%(xrange[0]+(xrange[1]-xrange[0])*k/4),
                 10,GREY,align=PP_ALIGN.CENTER)
    for xp,yp,cat,txt in points:
        xx=px+(xp-xrange[0])/(xrange[1]-xrange[0])*pw
        yy=py+ph-(yp-yrange[0])/(yrange[1]-yrange[0])*ph
        put_rect(sl,xx-.057,yy-.057,.114,.114,fill=colors[str(cat)],radius=True)
        if txt:put_text(sl,xx+.11,yy-.29,1.40,.23,txt,10,NAVY,bold=True)
    put_text(sl,x+w*.28,y+h-.16,w*.60,.21,"Peak die temperature [°C]",11,INK,align=PP_ALIGN.CENTER)

def sectionrule(sl,y=1.52):
    put_line(sl,.60,y,12.7,y,LINE,1.0)

def reference(sl,text):
    put_text(sl,.60,6.79,12.2,.18,text,size=8.2,color=GREY)
